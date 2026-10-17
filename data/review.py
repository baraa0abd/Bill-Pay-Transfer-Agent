"""Native-speaker review of the held-out set through a spreadsheet.

    python -m data.review export   -> data/heldout_review.xlsx   (edit in Excel, save)
    python -m data.review import   -> rewrites data/heldout.jsonl from your edits, validated against the wallet

Sheet 1 has one row per request; edit the white columns. Sheet 2 lists user 1's contacts, billers and balance.
Rows you mark 'keep' or 'edited' become hand_edited with your name; 'delete' removes the row.
"""
import json
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from .models import Seed, TestItem
from .synthetic import check_references, user_view

HERE = Path(__file__).parent
TESTSET, SEED, SHEET = HERE / 'heldout.jsonl', HERE / 'heldout_seed.json', HERE / 'heldout_review.xlsx'
COLS = ['id', 'text', 'outcome', 'ask_slot', 'payee_id', 'biller_id', 'amount_iqd', 'refuse_code', 'intent_types',
        'category', 'ai_note', 'decision', 'reviewer', 'your_note']
EDITABLE = {'text', 'outcome', 'ask_slot', 'payee_id', 'biller_id', 'amount_iqd', 'refuse_code', 'intent_types', 'decision', 'reviewer', 'your_note'}
CHOICES = {'outcome': 'EXECUTE_AFTER_CONFIRM,ASK,REFUSE,SPLIT', 'ask_slot': 'payee,biller,amount',
           'refuse_code': 'UNKNOWN_PAYEE,UNKNOWN_BILLER,INSUFFICIENT_FUNDS,LIMIT_EXCEEDED,BILL_ALREADY_PAID,AMOUNT_INVALID,UNSUPPORTED,UNCLEAR',
           'decision': 'keep,edited,delete'}


def export():
    items = [json.loads(l) for l in TESTSET.read_text(encoding='utf-8').splitlines()]
    wb = Workbook(); ws = wb.active; ws.title = 'requests'
    ws.append(COLS)
    for i in items:
        f = i['expected_final']
        note = i['human_review'].replace('Pending native Iraqi speaker review.', '').strip()
        ws.append([i['id'], i['text'], i['expected_outcome'], i.get('expected_ask_slot'), f.get('payee_id'), f.get('biller_id'),
                   f.get('amount_iqd'), f.get('refuse_code'), ','.join(i['expected_intent_types']), i['category'], note, None, None, None])
    grey = PatternFill('solid', fgColor='EEEEEE')
    for c, name in enumerate(COLS, 1):
        ws.cell(1, c).font = Font(bold=True)
        for row in ws.iter_rows(min_row=2, min_col=c, max_col=c):
            if name not in EDITABLE: row[0].fill = grey
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(horizontal='right', readingOrder=2, wrap_text=True)  # Arabic, right-to-left
    for name, values in CHOICES.items():
        dv = DataValidation(type='list', formula1=f'"{values}"', allow_blank=True)
        col = 'ABCDEFGHIJKLMN'[COLS.index(name)]; dv.add(f'{col}2:{col}{len(items) + 1}'); ws.add_data_validation(dv)
    for col, width in zip('ABCDEFGHIJKLMN', [12, 48, 22, 10, 9, 9, 11, 20, 20, 18, 55, 10, 14, 30]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'C2'
    # Reference sheet: what the labels refer to.
    wallet = user_view(Seed.model_validate_json(SEED.read_text(encoding='utf-8')))
    ref = wb.create_sheet('wallet (user 1)')
    ref.append(['balance_iqd', wallet['balance_iqd'], 'daily_limit_iqd', wallet['daily_limit_iqd'], 'fee per payment', wallet['fee_per_payment_iqd']])
    ref.append([]); ref.append(['contact id', 'name', 'nickname'])
    for c in wallet['contacts']: ref.append([c['id'], c['name'], c['nickname']])
    ref.append([]); ref.append(['biller id', 'name', 'category', 'unpaid bills (IQD)'])
    for b in wallet['billers']: ref.append([b['biller_id'], b['name'], b['category'], ', '.join(map(str, b['unpaid_bills_iqd'])) or 'none (fully paid)'])
    ref.column_dimensions['B'].width = 28
    wb.save(SHEET); print(f'Wrote {SHEET}. Edit it in Excel, save, then run: python -m data.review import')


def as_int(v):
    return int(v) if v not in (None, '') else None


def import_():
    rows = list(load_workbook(SHEET).active.iter_rows(min_row=2, values_only=True))
    old = {json.loads(l)['id']: json.loads(l) for l in TESTSET.read_text(encoding='utf-8').splitlines()}
    items, reviewed = [], 0
    for r in rows:
        r = dict(zip(COLS, r))
        if not r['id']: continue
        if (r['decision'] or '').strip() == 'delete': continue
        o = r['outcome']
        final = {'ASK': {}, 'SPLIT': {}, 'REFUSE': {'refuse_code': r['refuse_code']},
                 'EXECUTE_AFTER_CONFIRM': {k: v for k, v in dict(type='bill_payment' if r['biller_id'] else 'transfer', payee_id=as_int(r['payee_id']),
                                                                 biller_id=as_int(r['biller_id']), amount_iqd=as_int(r['amount_iqd'])).items() if v is not None}}[o]
        types = [t.strip() for t in str(r['intent_types']).split(',') if t.strip()]
        item = old[r['id']] | dict(text=str(r['text']).strip(), expected_outcome=o, expected_ask_slot=r['ask_slot'] if o == 'ASK' else None,
                                   expected_final=final, expected_intent_types=types, expected_intent_count=len(types))
        if r['decision'] in ('keep', 'edited') and r['reviewer']:
            reviewed += 1
            item |= dict(hand_edited=True, review_status='human_reviewed',
                         human_review=f"Reviewed by {r['reviewer']} ({r['decision']}). {r['your_note'] or ''}".strip())
        items.append(item)
    seed = Seed.model_validate_json(SEED.read_text(encoding='utf-8'))
    check_references([TestItem(**i) for i in items], seed)  # any typo in an ID or label fails here, before writing
    TESTSET.write_text(''.join(json.dumps(i, ensure_ascii=False) + '\n' for i in items), encoding='utf-8')
    print(f'Imported {len(items)} requests ({reviewed} reviewed by a person, {len(old) - len(items)} deleted) into {TESTSET.name}.')


if __name__ == '__main__':
    {'export': export, 'import': import_}[sys.argv[1]]()
